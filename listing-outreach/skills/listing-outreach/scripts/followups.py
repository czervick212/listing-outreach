#!/usr/bin/env python3
"""Follow-ups for outreach blasts — who was sent, who read it, who needs a nudge.

    python3 followups.py status [listing]        # sync with the mailbox, print where each blast stands
    python3 followups.py run [--dry-run]         # the daily job: sync, draft due nudges, notify
    python3 followups.py set <listing> [--followup "..."] [--followup2 "..."] [--cadence 3,10]
    python3 followups.py close <listing>         # stop following up on a blast (reopen: `reopen`)
    python3 followups.py import-legacy           # make a blast log from an old send-manifest.json
    python3 followups.py schedule install|remove|status

How a nudge happens. Every blast log (build_send_manifest.py writes it) carries one or two
follow-up messages written at blast time. `run`:
  1. reads recent mail both ways and folds it into each open blast — the original confirmed
     in Sent, read receipts, replies (lib/lo_blasts.py sync);
  2. for each recipient with no reply, drafts nudge 1 once 3 business days have passed since
     the original went out, and nudge 2 at 10 — only after nudge 1 was actually sent;
  3. drafts each nudge as a REPLY on the original thread (lib/lo_mail.py draft_followup) —
     never sends it;
  4. writes followups-due.md (and the copy named by `followups.due_note` in config, e.g. a
     note embedded in a daily note), and shows a desktop notification.

Nothing is ever sent. A reply at any point takes the recipient off the list for good.
"""
import datetime
import json
import os
import platform
import shutil
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
for cand in (os.path.join(HERE, "..", "..", "..", "lib"), os.path.join(HERE, "lib")):
    if os.path.isdir(cand):
        sys.path.insert(0, os.path.abspath(cand))
        break
import lo_blasts  # noqa: E402
import lo_mail  # noqa: E402

BASE = lo_blasts.BASE
CFG = os.path.join(BASE, "config.json")
DUE_MD = os.path.join(BASE, "followups-due.md")
LOG = os.path.join(BASE, "followups.log")
RUNTIME = os.path.join(BASE, "runtime")
LABEL = "com.listing-outreach.followups"
PLIST = os.path.expanduser(f"~/Library/LaunchAgents/{LABEL}.plist")
WIN_TASK = "ListingOutreachFollowups"


def cfg():
    try:
        with open(CFG) as f:
            return json.load(f)
    except Exception:
        return {}


def me():
    return ((cfg().get("user") or {}).get("email") or "").lower()


def log(msg):
    os.makedirs(BASE, exist_ok=True)
    with open(LOG, "a") as f:
        f.write(f"{lo_blasts.iso(lo_blasts.now())}  {msg}\n")


def short(dt_iso):
    d = lo_blasts.parse(dt_iso)
    return f"{d.month}/{d.day}" if d else ""


# ------------------------------------------------------------------ sync
def sync_all(blasts):
    """Read mail once, far enough back to cover the oldest open blast, and fold it in."""
    if not blasts:
        return {}
    import lo_scan
    oldest = min(lo_blasts.parse(b.get("built")) or lo_blasts.now() for b in blasts)
    days = min(max((lo_blasts.now() - oldest).days + 2, 3), 90)
    rows = lo_scan.gather(days=days, bodies=False)
    changes = {}
    for b in blasts:
        c = lo_blasts.sync(b, rows, me())
        if c:
            lo_blasts.save(b)
        changes[b["slug"]] = c
    return changes


def finished(b):
    """Nothing left to do: everyone replied, or had every nudge and a week to answer."""
    cadence = b.get("cadence") or lo_blasts.DEFAULT_CADENCE
    today = lo_blasts.now().date()
    for s in b["sends"]:
        if s.get("replied_at"):
            continue
        if not s.get("sent_at"):
            return False
        sent_fus = [f for f in s["followups"] if f.get("sent_at")]
        wanted = len([n for n in range(1, len(cadence) + 1) if (b.get("followup") or {}).get(str(n))])
        if len(sent_fus) < wanted:
            return False
        if sent_fus and lo_blasts.business_days_between(
                lo_blasts.parse(sent_fus[-1]["sent_at"]).date(), today) < 5:
            return False
    return True


# ------------------------------------------------------------------ reporting
def summary(b):
    n = len(b["sends"])
    counts = {}
    for s in b["sends"]:
        k = lo_blasts.state(s)
        counts[k] = counts.get(k, 0) + 1
    reads = sum(1 for s in b["sends"] if s.get("read_at"))
    parts = [f"{v} {k}" for k, v in sorted(counts.items(), key=lambda kv: -kv[1])]
    if b.get("receipts"):
        parts.append(f"{reads} read")
    return f"{n} sent to · " + " · ".join(parts)


def print_status(blasts):
    if not blasts:
        print("No outreach blasts on record yet — build_send_manifest.py creates one.")
        return
    for b in blasts:
        flag = "  [closed]" if b.get("closed") else ""
        print(f"\n{b['listing']}{flag}")
        print(f"  {summary(b)}")
        cad = b.get("cadence") or lo_blasts.DEFAULT_CADENCE
        msgs = b.get("followup") or {}
        print(f"  nudges at {' / '.join(map(str, cad))} business days · messages: "
              f"{', '.join('#' + k for k in sorted(msgs)) or 'none written'}"
              f" · receipts {'on' if b.get('receipts') else 'off'}")
        for s in sorted(b["sends"], key=lambda s: (lo_blasts.state(s), s.get("contact") or s["to"])):
            who = s.get("contact") or s["to"]
            tenants = ", ".join(t.get("tenant") or "" for t in s.get("tenants") or [])[:40]
            extra = []
            if s.get("sent_at"):
                extra.append(f"sent {short(s['sent_at'])}")
            if s.get("read_at"):
                extra.append(f"read {short(s['read_at'])}")
            if s.get("not_read_at"):
                extra.append(f"deleted unread {short(s['not_read_at'])}")
            for f in s["followups"]:
                extra.append(f"nudge {f['n']} " + (f"sent {short(f['sent_at'])}" if f.get("sent_at")
                                                   else f"drafted {short(f['drafted_at'])}"))
            if s.get("replied_at"):
                extra.append(f"replied {short(s['replied_at'])}")
            print(f"    {lo_blasts.state(s):<16} {who[:24]:<24} {tenants:<40} {' · '.join(extra)}")


def write_due(blasts, drafted):
    """followups-due.md — what was drafted today and where every open blast stands."""
    now = lo_blasts.now()
    lines = []
    total = sum(len(v) for v in drafted.values())
    if total:
        lines.append(f"**{total} follow-up{'s' if total != 1 else ''} drafted** — "
                     f"review and send from your drafts.")
    else:
        lines.append("No follow-ups due today.")
    lines.append("")
    for b in blasts:
        mine = drafted.get(b["slug"]) or []
        head = f"**{b['listing']}** — {summary(b)}"
        lines.append(f"- {head}")
        for s, n in mine:
            who = s.get("contact") or s["to"]
            read = f" · read {short(s['read_at'])}" if s.get("read_at") else ""
            lines.append(f"    - nudge {n} → {who} ({', '.join(t.get('tenant') or '' for t in s['tenants'])}){read}")
    lines.append("")
    lines.append(f"_Updated {now:%-m/%-d %-I:%M %p}_")
    text = "\n".join(lines) + "\n"
    os.makedirs(BASE, exist_ok=True)
    with open(DUE_MD, "w") as f:
        f.write(text)
    extra = (cfg().get("followups") or {}).get("due_note")
    if extra:
        try:
            os.makedirs(os.path.dirname(os.path.expanduser(extra)), exist_ok=True)
            with open(os.path.expanduser(extra), "w") as f:
                f.write(text)
        except Exception as e:
            log(f"couldn't write due_note {extra}: {e}")
    return total


def notify(title, message):
    try:
        if platform.system() == "Darwin":
            esc = lambda s: s.replace("\\", "\\\\").replace('"', '\\"')
            subprocess.run(["osascript", "-e",
                            f'display notification "{esc(message)}" with title "{esc(title)}"'],
                           capture_output=True, timeout=10)
        elif platform.system() == "Windows":
            ps = ("Add-Type -AssemblyName System.Windows.Forms;"
                  "$n=New-Object System.Windows.Forms.NotifyIcon;"
                  "$n.Icon=[System.Drawing.SystemIcons]::Information;$n.Visible=$true;"
                  f"$n.ShowBalloonTip(10000,'{title}','{message}',"
                  "[System.Windows.Forms.ToolTipIcon]::Info);Start-Sleep 11;$n.Dispose()")
            subprocess.Popen(["powershell", "-NoProfile", "-Command", ps])
    except Exception:
        pass


def update_sheet(b):
    """Best-effort: one narrow 'Tracking' column on the Targets sheet. Skipped when the
    workbook is missing or open in Excel (a save would be clobbered on their next save)."""
    path = b.get("sheet") or ""
    if not path or not os.path.exists(path):
        return False
    d, name = os.path.split(path)
    if os.path.exists(os.path.join(d, "~$" + name)):
        return False
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb["Targets"]
    except Exception:
        return False
    H = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    if "Email" not in H:
        return False
    col = H.get("Tracking") or ws.max_column + 1
    ws.cell(row=1, column=col).value = "Tracking"
    by = {s["to"].lower(): s for s in b["sends"]}
    for r in range(2, ws.max_row + 1):
        s = by.get(str(ws.cell(row=r, column=H["Email"]).value or "").strip().lower())
        if not s:
            continue
        bits = []
        if s.get("sent_at"):
            bits.append(f"Sent {short(s['sent_at'])}")
        if s.get("read_at"):
            bits.append(f"Read {short(s['read_at'])}")
        for f in s["followups"]:
            bits.append(f"Nudge {f['n']} {short(f.get('sent_at') or f['drafted_at'])}"
                        + ("" if f.get("sent_at") else " (draft)"))
        if s.get("replied_at"):
            bits.append(f"Replied {short(s['replied_at'])}")
        ws.cell(row=r, column=col).value = " · ".join(bits) or None
    try:
        wb.save(path)
        return True
    except Exception:
        return False


# ------------------------------------------------------------------ run
def run(dry=False, force=False):
    today = lo_blasts.now().date()
    if today.weekday() >= 5 and not force:
        print("Weekend — nothing drafted. (--force to run anyway)")
        return 0
    refresh_runtime()

    # A receipts switch left on past its deadline (the watcher died, the Mac slept) comes off here.
    if lo_mail.backend() == "apple-mail":
        try:
            import lo_receipts
            st = lo_receipts.state()
            if lo_receipts.is_on() and st and lo_blasts.parse(st.get("deadline")) and \
                    lo_blasts.now() > lo_blasts.parse(st["deadline"]):
                lo_receipts.off("deadline passed, caught by the daily run")
                log("receipts switch was still on past its deadline — turned off")
        except Exception:
            pass

    blasts = lo_blasts.open_blasts()
    if not blasts:
        write_due([], {})
        print("No open blasts.")
        return 0
    try:
        changes = sync_all(blasts)
    except Exception as e:
        log(f"sync failed: {e}")
        print(f"Couldn't read the mailbox: {e}", file=sys.stderr)
        return 2
    for slug, c in changes.items():
        for line in c:
            log(f"{slug}: {line}")

    drafted, failed = {}, []
    for b in blasts:
        for s, n in lo_blasts.due(b, today):
            text = lo_blasts.fill(b["followup"][str(n)], s, b)
            who = s.get("contact") or s["to"]
            if dry:
                print(f"  would draft nudge {n} → {who}\n      {text[:100]}")
                drafted.setdefault(b["slug"], []).append((s, n))
                continue
            ok, err = lo_mail.draft_followup(s, text, me())
            if ok:
                s["followups"].append({"n": n, "drafted_at": lo_blasts.iso(lo_blasts.now()),
                                       "sent_at": None})
                drafted.setdefault(b["slug"], []).append((s, n))
                log(f"{b['slug']}: drafted nudge {n} -> {s['to']}")
            else:
                failed.append(f"{b['listing']}: {who} ({err})")
                log(f"{b['slug']}: FAILED nudge {n} -> {s['to']}: {err}")
        if not dry:
            if finished(b):
                b["closed"] = True
                log(f"{b['slug']}: finished — closed")
            lo_blasts.save(b)
            update_sheet(b)

    if dry:
        print(f"\n(dry run — {sum(len(v) for v in drafted.values())} nudge(s) would be drafted)")
        return 0
    total = write_due(blasts, drafted)
    if total:
        where = ("your Outlook Drafts folder" if lo_mail.backend() == "graph"
                 else "compose windows in your mail app")
        names = ", ".join(b["listing"] for b in blasts if drafted.get(b["slug"]))
        notify("Outreach follow-ups", f"{total} drafted ({names}) — review in {where}.")
    print(f"{total} follow-up(s) drafted" + (f", {len(failed)} failed" if failed else ""))
    for f in failed:
        print(f"  FAILED {f}")
    return 0


# ------------------------------------------------------------------ scheduling
def refresh_runtime():
    """Keep a copy of this script + lib under ~/.listing-outreach/runtime for the scheduler.
    The plugin's own path changes with every version, and may sit somewhere a background
    job isn't allowed to read — the home folder always works."""
    src_lib = os.path.abspath(os.path.join(HERE, "..", "..", "..", "lib"))
    if not os.path.isdir(src_lib) or os.path.abspath(HERE) == os.path.abspath(RUNTIME):
        return
    os.makedirs(RUNTIME, exist_ok=True)
    shutil.copy2(os.path.abspath(__file__), os.path.join(RUNTIME, "followups.py"))
    dst = os.path.join(RUNTIME, "lib")
    os.makedirs(dst, exist_ok=True)
    for name in os.listdir(src_lib):
        if name.endswith(".py"):
            shutil.copy2(os.path.join(src_lib, name), os.path.join(dst, name))


def schedule(action):
    system = platform.system()
    target = os.path.join(RUNTIME, "followups.py")
    if action == "install":
        refresh_runtime()
        if system == "Darwin":
            # /usr/bin/python3 on purpose: it's the binary Full Disk Access is granted to for
            # reading ~/Library/Mail from a background job.
            plist = f"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE plist PUBLIC "-//Apple//DTD PLIST 1.0//EN" "http://www.apple.com/DTDs/PropertyList-1.0.dtd">
<plist version="1.0"><dict>
  <key>Label</key><string>{LABEL}</string>
  <key>ProgramArguments</key><array>
    <string>/usr/bin/python3</string><string>{target}</string><string>run</string>
  </array>
  <key>StartCalendarInterval</key><dict><key>Hour</key><integer>7</integer><key>Minute</key><integer>0</integer></dict>
  <key>StandardOutPath</key><string>{LOG}</string>
  <key>StandardErrorPath</key><string>{LOG}</string>
</dict></plist>
"""
            os.makedirs(os.path.dirname(PLIST), exist_ok=True)
            subprocess.run(["launchctl", "bootout", f"gui/{os.getuid()}", PLIST], capture_output=True)
            with open(PLIST, "w") as f:
                f.write(plist)
            p = subprocess.run(["launchctl", "bootstrap", f"gui/{os.getuid()}", PLIST],
                               capture_output=True, text=True)
            if p.returncode != 0:
                print(f"launchctl: {p.stderr.strip()}", file=sys.stderr)
                return 2
            print("Follow-ups scheduled: weekdays at 7:00 AM (runs when the Mac is awake).")
            return 0
        if system == "Windows":
            p = subprocess.run(["schtasks", "/Create", "/F", "/SC", "DAILY", "/ST", "07:00",
                                "/TN", WIN_TASK, "/TR", f'"{sys.executable}" "{target}" run'],
                               capture_output=True, text=True)
            print(p.stdout.strip() or p.stderr.strip())
            return p.returncode
        print("Scheduling is supported on macOS and Windows only.")
        return 1
    if action == "remove":
        if system == "Darwin":
            subprocess.run(["launchctl", "bootout", f"gui/{os.getuid()}", PLIST], capture_output=True)
            if os.path.exists(PLIST):
                os.remove(PLIST)
        elif system == "Windows":
            subprocess.run(["schtasks", "/Delete", "/F", "/TN", WIN_TASK], capture_output=True)
        print("Follow-up schedule removed.")
        return 0
    if system == "Darwin":
        on = subprocess.run(["launchctl", "print", f"gui/{os.getuid()}/{LABEL}"],
                            capture_output=True).returncode == 0
    elif system == "Windows":
        on = subprocess.run(["schtasks", "/Query", "/TN", WIN_TASK], capture_output=True).returncode == 0
    else:
        on = False
    print(f"Follow-up schedule: {'installed (weekdays 7:00 AM)' if on else 'not installed'}")
    return 0


# ------------------------------------------------------------------ misc commands
def import_legacy():
    try:
        with open(lo_blasts.LEGACY) as f:
            man = json.load(f)
    except Exception:
        print("No send-manifest.json to import.")
        return 1
    slug = lo_blasts.slugify(man.get("listing"))
    if os.path.exists(lo_blasts.path_for(slug)):
        print(f"{slug} already has a blast log.")
        return 0
    built = lo_blasts.parse(man.get("built")) or lo_blasts.now()
    blast = {"slug": slug, "listing": man.get("listing") or slug,
             "vts_property_id": man.get("vts_property_id"), "sheet": man.get("sheet"),
             "built": lo_blasts.iso(built), "receipts": False, "followup": {},
             "cadence": lo_blasts.DEFAULT_CADENCE, "closed": False,
             "sends": [lo_blasts.new_send(s["to"], s.get("contact", ""), s.get("firm", ""),
                                          s.get("subject") or "", s.get("body") or "",
                                          s.get("tenants")) for s in man.get("sends", [])]}
    lo_blasts.save(blast)
    print(f"Imported {blast['listing']} ({len(blast['sends'])} recipients). No follow-up "
          f"messages yet — add them with:  followups.py set \"{blast['listing']}\" --followup \"...\"")
    return 0


def main():
    args = sys.argv[1:]
    cmd = args[0] if args else "status"

    def opt(flag):
        return args[args.index(flag) + 1] if flag in args else None

    if cmd == "status":
        blasts = [lo_blasts.find(args[1])] if len(args) > 1 else lo_blasts.all_blasts()
        blasts = [b for b in blasts if b]
        if "--no-sync" not in args:
            try:
                sync_all([b for b in blasts if not b.get("closed")])
            except Exception as e:
                print(f"(couldn't read the mailbox, showing the last known state: {e})")
        print_status(blasts)
        return 0
    if cmd == "run":
        return run(dry="--dry-run" in args, force="--force" in args)
    if cmd in ("set", "close", "reopen"):
        if len(args) < 2 or not lo_blasts.find(args[1]):
            print(f"No single blast matches {args[1] if len(args) > 1 else '(none given)'}.")
            return 1
        b = lo_blasts.find(args[1])
        if cmd == "close":
            b["closed"] = True
        elif cmd == "reopen":
            b["closed"] = False
        else:
            b.setdefault("followup", {})
            if opt("--followup"):
                b["followup"]["1"] = opt("--followup")
            if opt("--followup2"):
                b["followup"]["2"] = opt("--followup2")
            if opt("--cadence"):
                b["cadence"] = [int(x) for x in opt("--cadence").split(",")]
        lo_blasts.save(b)
        print(f"{b['listing']}: saved.")
        return 0
    if cmd == "import-legacy":
        return import_legacy()
    if cmd == "schedule":
        return schedule(args[1] if len(args) > 1 else "status")
    print(__doc__)
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
