#!/usr/bin/env python3
"""Open the outreach emails as REVIEWABLE compose windows, flyer attached.

    python make_drafts.py <targets.xlsx> <flyer.pdf>            # all sends
    python make_drafts.py <targets.xlsx> <flyer.pdf> 0 10       # a batch (start, count)
    python make_drafts.py <targets.xlsx> <flyer.pdf> --list     # print the send list only
    python make_drafts.py <targets.xlsx> <flyer.pdf> --receipts [--blast "<Listing Name>"]
    python make_drafts.py <targets.xlsx> <flyer.pdf> --send [--wait 5]   # SEND, then bounce-check

--send (Apple Mail only) sends instead of drafting — only after the user has approved the list.
One at a time, each address logged to ~/.listing-outreach/sends/<sheet>.json as it goes, so a
re-run never sends twice. It then waits --wait minutes (default 5) for bounces, reads them from
Mail, and records delivered/bounced per address. Pass that log to build_vts_plan.py --send-log
so VTS is written only for the emails that landed.

--receipts asks each recipient's mail app for a read receipt. Outlook/Graph flag each draft.
Apple Mail can't do that per message, so it switches on Mail's receipt header for the whole
mailbox and starts a background watcher that switches it off once every email in the blast
is in Sent (6 hours at the latest). Build the blast log first (build_send_manifest.py) —
the watcher reads it; --blast names it, otherwise the most recent one is used.

(Use `python3` on macOS/Linux, `python` on Windows — try one, use the other if it's missing.)

Without --send nothing is sent — review each draft, then Send by hand. Where the drafts appear depends on
the backend (see lib/lo_mail.py): Outlook/Graph puts them in your Drafts folder, Apple Mail
and classic Outlook open them as compose windows on screen.
The From: address comes from ~/.listing-outreach/config.json (user.email); the signature is
whatever the mail client applies to that account.

Why script the client at all: mailto: links cannot carry an attachment. The flyer must live
somewhere normal — not inside the mail app's own container/Downloads — or attaching fails. Copy
it next to the listing first.
"""
import openpyxl
import sys
import os
import json

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "..", "lib"))
import lo_mail  # noqa: E402
import lo_blasts  # noqa: E402

if len(sys.argv) < 3:
    print(__doc__)
    raise SystemExit(1)

SHEET, FLYER = sys.argv[1], sys.argv[2]
rest = sys.argv[3:]
RECEIPTS = "--receipts" in rest
SEND = "--send" in rest
WAIT = 5
if "--wait" in rest:
    WAIT = float(rest[rest.index("--wait") + 1])
    rest = [a for i, a in enumerate(rest) if not (a == "--wait" or (i and rest[i - 1] == "--wait"))]
BLAST = rest[rest.index("--blast") + 1] if "--blast" in rest else None
if BLAST:
    rest = [a for a in rest if a not in ("--blast", BLAST)]

CFG = os.path.join(os.path.expanduser("~"), ".listing-outreach", "config.json")
try:
    SENDER = json.load(open(CFG))["user"]["email"]
except Exception:
    print("No sender email configured — run listing-outreach setup first.", file=sys.stderr)
    raise SystemExit(2)

if not os.path.exists(FLYER):
    print(f"Flyer not found: {FLYER}", file=sys.stderr)
    raise SystemExit(1)

ws = openpyxl.load_workbook(SHEET)["Targets"]
H = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
SEND_COL = H["Send Email"]

sends = []
for r in range(2, ws.max_row + 1):
    if ws.cell(row=r, column=SEND_COL).hyperlink:          # only rows with a live Send button
        sends.append({
            "tenant": ws.cell(row=r, column=H["Tenant"]).value,
            "contact": ws.cell(row=r, column=H["Broker / Contact"]).value,
            "email": ws.cell(row=r, column=H["Email"]).value,
            "subject": ws.cell(row=r, column=H["Subject"]).value,
            "body": ws.cell(row=r, column=H["Email Draft"]).value,
        })

if "--list" in rest:
    for i, s in enumerate(sends):
        print(f"{i:>3}  {s['contact'] or '-':<20} {s['email']:<34} {s['tenant']}")
    print(f"\n{len(sends)} sends")
    raise SystemExit

nums = [a for a in rest if a.isdigit()]
START = int(nums[0]) if len(nums) > 0 else 0
COUNT = int(nums[1]) if len(nums) > 1 else len(sends)
batch = sends[START:START + COUNT]
print(f"total sends: {len(sends)} | drafting {len(batch)} "
      f"(index {START}-{START + len(batch) - 1})")

if RECEIPTS and lo_mail.backend() == "apple-mail":
    import lo_receipts
    if BLAST:
        slug = lo_blasts.slugify(BLAST)
    else:
        try:
            slug = lo_blasts.slugify(json.load(open(lo_blasts.LEGACY))["listing"])
        except Exception:
            slug = None
    if not slug or not os.path.exists(lo_blasts.path_for(slug)):
        print("--receipts needs the blast log first: run build_send_manifest.py, then pass "
              "--blast \"<Listing Name>\".", file=sys.stderr)
        raise SystemExit(2)
    deadline = lo_receipts.on(slug, SENDER)
    lo_receipts.start_watcher(slug)
    print(f"Read receipts ON in Apple Mail until this blast is sent (auto-off by "
          f"{deadline:%-I:%M %p}). Anything else you send before then asks for one too.")

# Keep the flyer with the blast log: every nudge re-attaches it (follow-ups.py).
_slug = lo_blasts.slugify(BLAST) if BLAST else None
if not _slug:
    try:
        _slug = lo_blasts.slugify(json.load(open(lo_blasts.LEGACY))["listing"])
    except Exception:
        _slug = None
if _slug and os.path.exists(lo_blasts.path_for(_slug)):
    _b = json.load(open(lo_blasts.path_for(_slug)))
    _kept = lo_blasts.keep_flyer(_slug, os.path.abspath(FLYER))
    if _kept and _b.get("flyer") != _kept:
        _b["flyer"] = _kept
        lo_blasts.save(_b)

if SEND:
    import time
    import lo_bounces
    os.makedirs(os.path.join(lo_blasts.BASE, "sends"), exist_ok=True)
    LOG = os.path.join(lo_blasts.BASE, "sends",
                       os.path.splitext(os.path.basename(SHEET))[0] + ".json")
    started = time.time() - 60
    sent, skipped, failed = lo_mail.send_now(batch, os.path.abspath(FLYER), SENDER, LOG)
    print(f"\nsent {sent}, already sent {skipped}, failed {failed}  (log: {LOG})")
    if sent and WAIT > 0:
        print(f"waiting {WAIT:g} min for bounces before anything is logged to VTS...")
        time.sleep(WAIT * 60)
    log = json.load(open(LOG)) if os.path.exists(LOG) else {}
    mine = [k for k, v in log.items() if v.get("sent")]
    since = min([log[k]["at"] for k in mine] + [started]) - 60
    bad = lo_bounces.bounced(mine, since)
    if bad is None:
        print("couldn't read Mail's index — bounce check skipped; check Mail before logging VTS.")
    else:
        for k in mine:
            log[k]["bounced"] = k in bad
        with open(LOG, "w") as f:
            json.dump(log, f, indent=1)
        landed = [k for k in mine if k not in bad]
        print(f"delivered {len(landed)}, bounced {len(bad)}")
        for k in sorted(bad):
            print(f"  BOUNCED  {log[k]['tenant']} <{k}> — needs a new contact; no VTS comment")
    raise SystemExit(1 if failed else 0)

ok, fail = lo_mail.open_drafts(batch, os.path.abspath(FLYER), SENDER, receipts=RECEIPTS)
where = ("your Outlook Drafts folder" if lo_mail.backend() == "graph"
         else "the compose windows on screen")
print(f"\ndrafted {ok}, failed {fail} — review in {where} and send by hand. "
      f"Nothing was sent automatically.")
